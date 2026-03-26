import openpyxl

# Full analysis of codici ateco-settori
print("=== CODICI ATECO -> SETTORE (complete mapping) ===")
wb1 = openpyxl.load_workbook(r"C:\Users\mcollivignarelli\Downloads\codici ateco-settori.xlsx", data_only=True)
ws1 = wb1["Foglio1"]
mapping_sectors = set()
for r in range(1, ws1.max_row + 1):
    code = str(ws1.cell(r, 1).value or "")
    sector = str(ws1.cell(r, 2).value or "")
    mapping_sectors.add(sector)
    print(f"  '{code}' → '{sector}'")

print(f"\nUnique sectors: {sorted(mapping_sectors)}")

# Check the Clienti Leyton data for region distribution
print("\n=== CLIENTI LEYTON - Region + Sub-industry analysis ===")
wb2 = openpyxl.load_workbook(r"C:\Users\mcollivignarelli\Downloads\Clienti Leyton.xlsx", data_only=True)
ws2 = wb2["report1774361668537"]

from collections import Counter
regions = Counter()
sub_industries = Counter()
products = Counter()
no_region = 0
total = 0

for r in range(2, ws2.max_row + 1):
    name = str(ws2.cell(r, 1).value or "").upper()
    if "SEGNALATORE" in name or "PARTNERSHIP" in name:
        continue
    total += 1
    
    region = str(ws2.cell(r, 11).value or "None")
    if region == "None" or region.strip() == "":
        no_region += 1
    else:
        regions[region] += 1
    
    sub_ind = str(ws2.cell(r, 8).value or "")
    if sub_ind and sub_ind != "None":
        # Extract 2-digit code
        code_2d = sub_ind.split(" ")[0].strip() if sub_ind else ""
        sub_industries[code_2d] += 1
    
    prods = str(ws2.cell(r, 10).value or "")
    if prods and prods != "None":
        for p in prods.split(";"):
            p = p.strip()
            if p:
                products[p] += 1

print(f"\nTotal valid rows: {total}")
print(f"Rows without region: {no_region}")
print(f"\nTop 20 regions:")
for reg, cnt in regions.most_common(20):
    print(f"  {reg}: {cnt}")

print(f"\nTop 20 sub-industry codes:")
for code, cnt in sub_industries.most_common(20):
    print(f"  {code}: {cnt}")

print(f"\nTop 20 products:")
for p, cnt in products.most_common(20):
    print(f"  {p}: {cnt}")
